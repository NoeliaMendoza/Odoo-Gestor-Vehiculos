"""Crea documentos Word editables a partir de las mismas secciones del HTML."""

import base64

from io import BytesIO



from docx import Document

from docx.shared import Cm, Pt, RGBColor

from docx.oxml import OxmlElement

from docx.oxml.ns import qn

from odoo import tools





def commitment_date(value):

    """Presenta la fecha local del acta en español, sin etiquetas técnicas."""

    day, month, year = value.split('/')

    months = ('enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio',

              'julio', 'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre')

    return f'Santo Domingo, {int(day)} de {months[int(month) - 1]} de {year}'





def commitment_template():

    """Abre una copia en memoria del Word original entregado por la empresa."""

    with tools.file_open('control_vehiculos/data/commitment_template.docx', 'rb') as source:

        return Document(source)





def commitment_paragraphs():

    """Conserva el texto, las negritas y las listas del original en la vista HTML."""

    return [{'center': paragraph.alignment == 1, 'listed': paragraph.style.name == 'List Paragraph',

             'runs': [{'text': run.text, 'bold': bool(run.bold),

                       'size': run.font.size.pt if run.font.size else 12} for run in paragraph.runs if run.text]}

            for paragraph in commitment_template().paragraphs if paragraph.text]





def build_commitment_word(data):

    """Sitúa fecha y firmas al pie de la hoja, reservando espacio para no tapar el texto."""

    doc = commitment_template()

    page = doc.sections[0]

    page.bottom_margin = Cm(4.2)

    page.footer_distance = Cm(1.4)

    footer = page.footer

    paragraph = footer.paragraphs[0]

    paragraph.alignment = 2

    paragraph.paragraph_format.space_after = Pt(28)

    run = paragraph.add_run(data['date_label'])

    run.font.size = Pt(11)

    signatures = footer.add_table(rows=1, cols=2,

                                  width=page.page_width - page.left_margin - page.right_margin)

    for cell, (role, name) in zip(signatures.rows[0].cells, data['signatures']):

        cell.text = '____________________________\n' + name + '\n' + role

        p = cell.paragraphs[0]

        p.alignment = 1

        p.paragraph_format.space_after = Pt(0)

        p.paragraph_format.line_spacing = 1

        for run in p.runs:

            run.font.size = Pt(10)

    node = OxmlElement('w:cantSplit'); signatures.rows[0]._tr.get_or_add_trPr().append(node)

    # Word exige un párrafo después de la tabla; mantenerlo pequeño evita desplazar las firmas.

    tail = footer.add_paragraph()

    tail.paragraph_format.space_before = tail.paragraph_format.space_after = Pt(0)

    tail.paragraph_format.line_spacing = Pt(1)

    tail.add_run().font.size = Pt(1)

    stream = BytesIO(); doc.save(stream)

    return stream.getvalue()





def build_word(data):

    """Devuelve un DOCX real, con tablas legibles y espacios para firma manuscrita."""

    if data['kind'] == 'compromiso':

        return build_commitment_word(data)

    doc = Document()

    page = doc.sections[0]

    page.page_width, page.page_height = Cm(21), Cm(29.7)

    page.top_margin = page.bottom_margin = Cm(1.6)

    page.left_margin = page.right_margin = Cm(1.8)

    normal = doc.styles['Normal']

    normal.font.name, normal.font.size = 'Calibri', Pt(10)

    normal.paragraph_format.space_after = Pt(5)

    for name in ('Title', 'Heading 1', 'Heading 2'):

        doc.styles[name].font.color.rgb = RGBColor.from_string('243746')

    doc.styles['Title'].font.size = Pt(22)

    doc.styles['Heading 2'].font.size = Pt(11)

    doc.styles['Heading 2'].paragraph_format.space_before = Pt(12)

    header = page.header.paragraphs[0]

    header.add_run().add_picture(BytesIO(base64.b64decode(data['logo'])), width=Cm(4.8))

    header.add_run('   CONTROL DE VEHÍCULOS').font.size = Pt(9)

    footer = page.footer.paragraphs[0]

    footer.text = data['folio'] + ' · Documento de control interno   |   Página '

    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)

    doc.add_heading(data['title'], 0)

    doc.add_paragraph(data['subtitle'])

    p = doc.add_paragraph()

    p.add_run(data['folio']).bold = True

    p.add_run('   |   Estado: ' + data['state'])

    p = doc.add_paragraph(f'Emitido: {data["generated"]} · Hora: {data["timezone"]}')

    p.runs[0].font.size = Pt(9)

    for alert in data['alerts']:

        p = doc.add_paragraph(alert); p.runs[0].bold = True



    def table(headers, rows):

        if not rows:

            return

        t = doc.add_table(rows=0, cols=len(rows[0]))

        t.style = 'Table Grid'

        if headers:

            cells = t.add_row().cells

            for cell, text in zip(cells, headers):

                cell.text = text

                cell.paragraphs[0].runs[0].bold = True

                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'EEF2F5'); cell._tc.get_or_add_tcPr().append(shade)

            repeat = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)

        for row in rows:

            cells = t.add_row().cells

            for i, (cell, text) in enumerate(zip(cells, row)):

                cell.text = str(text or '—')

                if not headers and i == 0:

                    cell.paragraphs[0].runs[0].bold = True

        for row in t.rows:

            # Las filas normales se mantienen juntas; el texto excepcionalmente largo puede fluir.

            if sum(len(c.text) for c in row.cells) < 800:

                node = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(node)

        doc.add_paragraph().paragraph_format.space_after = Pt(0)



    for section in data['sections']:

        doc.add_heading(section['title'], level=2)

        if section['headers']:

            table(section['headers'], section['rows'])

        elif section['rows']:

            # Ficha de datos sin cuadrícula: etiqueta arriba y valor debajo.

            details = doc.add_table(rows=0, cols=2)

            for index in range(0, len(section['rows']), 2):

                cells = details.add_row().cells

                for cell, (label, value) in zip(cells, section['rows'][index:index+2]):

                    p = cell.paragraphs[0]

                    p.paragraph_format.space_after = Pt(8)

                    p.add_run(label + '\n').bold = True

                    p.add_run(str(value or '—'))

                node = OxmlElement('w:cantSplit'); details.rows[-1]._tr.get_or_add_trPr().append(node)

        if section['gauges']:

            gauges = doc.add_table(rows=1, cols=len(section['gauges']))

            for cell, gauge in zip(gauges.rows[0].cells, section['gauges']):

                cell.paragraphs[0].add_run(gauge['label']).bold = True

                cell.add_paragraph().add_run().add_picture(BytesIO(base64.b64decode(gauge['image'])), width=Cm(7.2))

                if gauge['low']:

                    cell.add_paragraph('Por debajo del mínimo requerido.').runs[0].bold = True

            node = OxmlElement('w:cantSplit'); gauges.rows[0]._tr.get_or_add_trPr().append(node)

        for text in section['paragraphs']:

            doc.add_paragraph(text)



    doc.add_heading('Firmas de constancia', level=2)

    p = doc.add_paragraph(data['statement']); p.paragraph_format.keep_with_next = True

    signatures = doc.add_table(rows=1, cols=2)

    for cell, (role, name) in zip(signatures.rows[0].cells, data['signatures']):

        cell.text = '\n\n____________________________\n' + name + '\n' + role + '\nFecha: _____________________'

    node = OxmlElement('w:cantSplit'); signatures.rows[0]._tr.get_or_add_trPr().append(node)

    stream = BytesIO(); doc.save(stream)

    return stream.getvalue()
